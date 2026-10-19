import openpyxl

wb = openpyxl.load_workbook('Job Fair Master Sheet.xlsx')
# Check each company sheet for disability-related columns
for sname in wb.sheetnames:
    if sname == 'JDs Overview ':
        continue
    ws = wb[sname]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    # find disability-related headers
    dis_cols = [(i, h) for i, h in enumerate(headers) if h and any(w in str(h).lower() for w in ['disab', 'pwd', 'type of', 'benchmark', 'eligible'])]
    if dis_cols:
        print(f"\n=== {sname} ===")
        print(f"  Headers: {headers}")
        print(f"  Disability cols: {dis_cols}")
        # show first 3 data rows for disability cols
        for row in list(ws.iter_rows(min_row=2, max_row=4)):
            vals = {i: str(row[i].value or '')[:60] for i, _ in dis_cols}
            print(f"  Row: {vals}")
    else:
        print(f"\n{sname}: No disability-specific cols. Headers: {[h for h in headers if h]}")
